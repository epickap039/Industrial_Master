"""API router: mrp."""
import ast
import io
import json
import math
import os
import re
import shutil
import socket
import subprocess
import sys
import traceback
import uuid
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

import openpyxl
import pandas as pd
import pyodbc
from fastapi import APIRouter, BackgroundTasks, File, Form, Header, HTTPException, Request, Response, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from database import get_db_connection, _int_from_count_row
from models import *

router = APIRouter()

# ─── Tabla de conversión fracción / calibre → mm canónico ──────────────────
# Permite derivar el espesor estándar directamente del nombre del material
# (p. ej. "ACERO ASTM A36 1/8"" → 3.18 mm) sin depender del valor crudo de
# Espesor_Perfil_CAD, que varía pieza a pieza y causa fragmentación en MRP.
_FRAC_TO_MM: List[tuple] = [
    ('1/16"', 1.59), ('1/8"',  3.18), ('3/16"', 4.76), ('1/4"',  6.35),
    ('5/16"', 7.94), ('3/8"',  9.53), ('7/16"', 11.11), ('1/2"', 12.70),
    ('9/16"', 14.29), ('5/8"', 15.88), ('3/4"', 19.05), ('7/8"', 22.23),
    ('1"',   25.40), ('1 1/4"', 31.75), ('1 1/2"', 38.10),
]
_CAL_TO_MM: List[tuple] = [
    ('C.10', 3.43), ('CAL.10', 3.43), ('CAL 10', 3.43),
    ('C.11', 3.04), ('CAL.11', 3.04), ('CAL 11', 3.04),
    ('C.14', 1.90), ('CAL.14', 1.90), ('CAL 14', 1.90),
    ('C.16', 1.52), ('CAL.16', 1.52), ('CAL 16', 1.52),
    ('C.18', 1.21), ('CAL.18', 1.21),
    ('C.19', 1.05), ('CAL.19', 1.05),
    ('C.20', 0.91), ('CAL.20', 0.91),
]


def _calibre_canonico(material: str) -> str:
    """Extrae el calibre/espesor canónico del nombre de material oficial.

    Prioriza fracciones de pulgada (p. ej. 1/8") sobre calibres GA (C.10…).
    Si no encuentra token conocido devuelve 'N/A'.
    """
    upper = material.upper()
    for token, mm in _FRAC_TO_MM:
        if token.upper() in upper:
            return f"{mm:.2f} mm ({token})"
    for token, mm in _CAL_TO_MM:
        if token in upper:
            return f"{mm:.2f} mm ({token})"
    return "N/A"


# ── Configuración manual de sugerencias de compra (por revisión) ─────────────
_MRP_COMPRA_CONFIG_PATH = os.path.normpath(
    os.path.join(os.path.dirname(__file__), "..", "data", "mrp_compra_config.json")
)
_MRP_MATERIAL_DEFAULTS_PATH = os.path.normpath(
    os.path.join(os.path.dirname(__file__), "..", "data", "mrp_material_defaults.json")
)

FORMATOS_COMPRA_MP: List[Dict[str, Any]] = [
    {"id": "auto", "tipo": "any", "etiqueta": "Automático (del nombre)", "area_m2": None, "longitud_m": None},
    {
        "id": "4x10", "tipo": "placa", "etiqueta": "4'×10'", "area_m2": 3.716,
        "largo_pies": 10.0, "ancho_pies": 4.0, "match_tokens": [],
    },
    {
        "id": "4x8", "tipo": "placa", "etiqueta": "4'×8'", "area_m2": 2.973,
        "largo_pies": 8.0, "ancho_pies": 4.0, "match_tokens": ["4X8", "4'X8", "4' X 8"],
    },
    {
        "id": "4x20", "tipo": "placa", "etiqueta": "4'×20'", "area_m2": 7.432,
        "largo_pies": 20.0, "ancho_pies": 4.0, "match_tokens": ["4X20", "4'X20", "4' X 20"],
    },
    {
        "id": "4x40", "tipo": "placa", "etiqueta": "4'×40'", "area_m2": 14.864,
        "largo_pies": 40.0, "ancho_pies": 4.0, "match_tokens": ["4X40", "4'X40", "4' X 40"],
    },
    {
        "id": "5x20", "tipo": "placa", "etiqueta": "5'×20'", "area_m2": 9.290,
        "largo_pies": 20.0, "ancho_pies": 5.0, "match_tokens": ["5'X20'", "5X20", "5' X 20"],
    },
    {
        "id": "5x24", "tipo": "placa", "etiqueta": "5'×24'", "area_m2": 11.148,
        "largo_pies": 24.0, "ancho_pies": 5.0, "match_tokens": ["5'X24'", "5X24"],
    },
    {
        "id": "8x20", "tipo": "placa", "etiqueta": "8'×20'", "area_m2": 14.864,
        "largo_pies": 20.0, "ancho_pies": 8.0, "match_tokens": ["8'X20'", "8X20"],
    },
    {
        "id": "8x30", "tipo": "placa", "etiqueta": "8'×30'", "area_m2": 22.297,
        "largo_pies": 30.0, "ancho_pies": 8.0, "match_tokens": ["8'X30'", "8X30"],
    },
    {
        "id": "tramo_5m", "tipo": "perfil", "etiqueta": "5 MT",
        "longitud_m": 5.0, "distancia_metros": 5.0, "solo_hss": False,
    },
    {
        "id": "tramo_5_8m", "tipo": "perfil", "etiqueta": "5.8 MT (Al)",
        "longitud_m": 5.8, "distancia_metros": 5.8, "solo_hss": False,
    },
    {
        "id": "tramo_7m", "tipo": "perfil", "etiqueta": "7 MT",
        "longitud_m": 7.0, "distancia_metros": 7.0, "solo_hss": False,
    },
    {
        "id": "tramo_6m", "tipo": "perfil", "etiqueta": "6 MT",
        "longitud_m": 6.0, "distancia_metros": 6.0, "solo_hss": False,
    },
    {
        "id": "hss_12m", "tipo": "perfil", "etiqueta": "12 MT (HSS)",
        "longitud_m": 12.0, "distancia_metros": 12.0, "solo_hss": True,
    },
    {"id": "manual", "tipo": "manual", "etiqueta": "Texto personalizado", "area_m2": None, "longitud_m": None},
    {
        "id": "directa", "tipo": "directa",
        "etiqueta": "Compra directa (por cantidad)",
        "area_m2": None, "longitud_m": None,
    },
]

_SCRAP_COMPRA_DEFAULT = 1.15
_FT2_TO_M2 = 0.09290304  # 1 pie² = 0.09290304 m²

# Expresiones regulares para extraer medidas del nombre del material.
# Se compilan una sola vez para eficiencia.
import re as _re

# Captura pies en el nombre: "4' X 10'", "4'X10'", "4 X 10 CAL", "4X8", etc.
# El ' al final es opcional porque a veces solo se escribe "4 X 10 CAL 3/8".
_RE_PLACA_PIES = _re.compile(
    r"""
    (\d+(?:\.\d+)?)\s*'?\s*[Xx×]\s*(\d+(?:\.\d+)?)\s*(?:'|PIES|FT|(?=\s+CAL|\s+III|\s+$|\s+ASTM|\s+AISI))
    """,
    _re.VERBOSE | _re.IGNORECASE,
)
# Captura metros en el nombre: "A 12 MT", "A 6 MT", "6 M ", "12MT"
_RE_METROS = _re.compile(
    r"""
    (?:A\s+|X\s+)?(\d+(?:\.\d+)?)\s*M(?:T|TS)?\b
    """,
    _re.VERBOSE | _re.IGNORECASE,
)


def _material_tipo_compra(material_upper: str) -> str:
    """Auto-clasificación por palabras clave.  Puede ser sobreescrita por tipo_compra en cfg."""
    if any(
        x in material_upper
        for x in (
            "PERFIL", "TUBO", "BARRA", "SOLERA", "ANGULO", "CANAL", "HSS",
            "REDONDO", "REDOND", "ROUND", "PTR", "IPR",
            "SPRING", "RESORTE", "RIEL",
        )
    ):
        return "perfil"
    if any(x in material_upper for x in ("PLACA", "LAMINA", "LÁMINA", "SHEET")):
        return "placa"
    return "placa"


def _pies_a_area_m2(largo_pies: float, ancho_pies: float) -> float:
    if largo_pies <= 0 or ancho_pies <= 0:
        return 0.0
    return float(largo_pies) * float(ancho_pies) * _FT2_TO_M2


def _float_cfg(v) -> Optional[float]:
    if v is None or v == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _normalize_cfg_entry(v: Dict[str, Any]) -> Dict[str, Any]:
    if not isinstance(v, dict):
        return {}
    tipo_raw = (v.get("tipo_compra") or "").strip().lower()
    return {
        "habilitado": v.get("habilitado", True) is not False,
        "formato_id": (v.get("formato_id") or "auto").strip(),
        # "tipo_compra": override manual ("placa" | "perfil" | "" = auto)
        "tipo_compra": tipo_raw if tipo_raw in ("placa", "perfil") else "",
        "largo_pies": _float_cfg(v.get("largo_pies")),
        "ancho_pies": _float_cfg(v.get("ancho_pies")),
        "distancia_metros": _float_cfg(v.get("distancia_metros")),
        "texto": (v.get("texto") or "").strip(),
    }


def _formato_compra_por_id(fmt_id: Optional[str]) -> Optional[Dict[str, Any]]:
    if not fmt_id:
        return None
    for f in FORMATOS_COMPRA_MP:
        if f["id"] == fmt_id:
            return f
    return None


def _detectar_formato_placa(material_upper: str) -> Dict[str, Any]:
    """Detecta formato de placa: primero por tokens fijos, luego regex desde el nombre."""
    for f in FORMATOS_COMPRA_MP:
        if f.get("tipo") != "placa":
            continue
        for tok in f.get("match_tokens") or []:
            if tok.upper() in material_upper:
                return f
    # Intentar extraer dimensiones en pies directamente del nombre
    m = _RE_PLACA_PIES.search(material_upper)
    if m:
        a_pies = float(m.group(1))
        l_pies = float(m.group(2))
        if a_pies > l_pies:
            a_pies, l_pies = l_pies, a_pies  # ancho ≤ largo
        area = _pies_a_area_m2(l_pies, a_pies)
        etiq = f"{int(a_pies) if a_pies == int(a_pies) else a_pies}'"
        etiq += f"×{int(l_pies) if l_pies == int(l_pies) else l_pies}'"
        return {
            "id": "auto_parsed",
            "tipo": "placa",
            "etiqueta": etiq,
            "area_m2": area,
            "largo_pies": l_pies,
            "ancho_pies": a_pies,
            "match_tokens": [],
        }
    return _formato_compra_por_id("4x10") or FORMATOS_COMPRA_MP[0]


def _detectar_formato_perfil(material_upper: str) -> Dict[str, Any]:
    """Detecta longitud de tramo: primero regex desde el nombre, luego HSS/default."""
    m = _RE_METROS.search(material_upper)
    if m:
        metros = float(m.group(1))
        if 1.0 <= metros <= 20.0:  # rango razonable
            etiq = f"{int(metros) if metros == int(metros) else metros} MT"
            return {
                "id": "auto_parsed",
                "tipo": "perfil",
                "etiqueta": etiq,
                "longitud_m": metros,
                "distancia_metros": metros,
            }
    if "HSS" in material_upper:
        return _formato_compra_por_id("hss_12m") or FORMATOS_COMPRA_MP[8]
    return _formato_compra_por_id("tramo_6m") or FORMATOS_COMPRA_MP[7]


def _cantidad_compra_auto(
    tipo: str,
    formato: Dict[str, Any],
    req_area_mm2: float,
    req_long_mm: float,
    scrap: float,
) -> float:
    """Devuelve cantidad exacta con decimales (sin techo) para mostrar el aprovechamiento real."""
    if tipo == "perfil":
        metros = req_long_mm / 1000.0
        tramos = float(formato.get("longitud_m") or 6.0)
        if tramos <= 0:
            tramos = 6.0
        return max(0.0, (metros * scrap) / tramos)
    area_m2 = float(formato.get("area_m2") or 3.716)
    m2_totales = req_area_mm2 / 1_000_000.0
    return max(0.0, (m2_totales * scrap) / area_m2)


def _fmt_qty(qty: float) -> str:
    """Formatea cantidad decimal: sin decimales si es entero, hasta 2 dígitos si no."""
    if qty <= 0:
        return "0"
    r = round(qty, 2)
    if r == int(r):
        return str(int(r))
    s = f"{r:.2f}".rstrip('0')
    return s if not s.endswith('.') else s[:-1]


def _texto_sugerencia_compra(
    tipo: str, formato: Dict[str, Any], cantidad: float, material: str = ""
) -> str:
    if formato.get("tipo") == "directa":
        return f"Comprar {math.ceil(max(0.0, cantidad))} {material}".strip() if cantidad > 0 else "Sin requerimiento"
    if cantidad <= 0:
        return "Sin requerimiento de compra"
    if formato.get("id") == "manual":
        return (material or "").strip() or "—"
    qty_str = _fmt_qty(cantidad)
    if tipo == "perfil":
        lm = float(formato.get("longitud_m") or 6)
        lm_str = str(int(lm)) if lm == int(lm) else str(lm)
        return f"Comprar {qty_str} Tramos de {lm_str} MT"
    etiq = formato.get("etiqueta") or "4'×10'"
    return f"Comprar {qty_str} Placas de {etiq}"


# Códigos de pieza (no materiales) que siempre se compran de forma directa por cantidad.
# Se excluyen del cálculo de área/perfil y del listado de auditoría CAD.
_DIRECTA_CODIGOS: frozenset = frozenset({
    "JE-012",
})


def _directa_codigos_sql_in() -> str:
    """Genera el literal SQL '(...)' para usar en cláusulas IN.
    Solo para listas de confianza definidas en el código (no entrada del usuario)."""
    if not _DIRECTA_CODIGOS:
        return "('__NEVER_MATCH__')"
    safe = sorted(c.replace("'", "") for c in _DIRECTA_CODIGOS)
    return "(" + ", ".join(f"'{c}'" for c in safe) + ")"


def _sugerencia_compra_automatica(
    material_oficial: str,
    req_area_mm2: float,
    req_long_mm: float,
    scrap: float = _SCRAP_COMPRA_DEFAULT,
) -> str:
    material_upper = material_oficial.upper()
    # Subensambles / compra directa por cantidad (no requieren cálculo de área)
    _DIRECTA_TOKENS = (
        "SEGURO DE RESORTE",
        "SEGURO RESORTE",
        "SPRING SEAL",
        "SPRING CLIP",
        "SUBENSAMBLE",
    )
    if any(tok in material_upper for tok in _DIRECTA_TOKENS):
        return "Compra directa por cantidad"
    if material_upper in {c.upper() for c in _DIRECTA_CODIGOS}:
        return "Compra directa por cantidad"
    if req_area_mm2 <= 0 and req_long_mm <= 0:
        return "Pendiente: cargar dimensiones CAD/DXF"
    tipo = _material_tipo_compra(material_upper)
    if tipo == "perfil":
        fmt = _detectar_formato_perfil(material_upper)
    else:
        fmt = _detectar_formato_placa(material_upper)
    qty = _cantidad_compra_auto(tipo, fmt, req_area_mm2, req_long_mm, scrap)
    return _texto_sugerencia_compra(tipo, fmt, qty)


def _formato_desde_cfg(
    cfg: Dict[str, Any],
    tipo: str,
    material_upper: str,
) -> Dict[str, Any]:
    """Arma el formato efectivo usando dropdown + medidas guardadas (pies / metros)."""
    fmt_id = (cfg.get("formato_id") or "auto").strip()
    if fmt_id == "auto":
        base = (
            _detectar_formato_perfil(material_upper)
            if tipo == "perfil"
            else _detectar_formato_placa(material_upper)
        )
    else:
        base = _formato_compra_por_id(fmt_id) or (
            _detectar_formato_perfil(material_upper)
            if tipo == "perfil"
            else _detectar_formato_placa(material_upper)
        )

    if tipo == "placa":
        lp = _float_cfg(cfg.get("largo_pies"))
        ap = _float_cfg(cfg.get("ancho_pies"))
        if lp is None and base.get("largo_pies") is not None:
            lp = float(base["largo_pies"])
        if ap is None and base.get("ancho_pies") is not None:
            ap = float(base["ancho_pies"])
        if lp and ap and lp > 0 and ap > 0:
            area_m2 = _pies_a_area_m2(lp, ap)
            etiq = f"{lp:g}'×{ap:g}'"
            return {
                **base,
                "id": base.get("id", "custom"),
                "etiqueta": etiq,
                "area_m2": area_m2,
                "largo_pies": lp,
                "ancho_pies": ap,
            }
        return base

    dist = _float_cfg(cfg.get("distancia_metros"))
    if dist is None and base.get("distancia_metros") is not None:
        dist = float(base["distancia_metros"])
    if dist is None and base.get("longitud_m") is not None:
        dist = float(base["longitud_m"])
    if dist and dist > 0:
        lm = float(dist)
        return {
            **base,
            "id": base.get("id", "custom"),
            "longitud_m": lm,
            "distancia_metros": lm,
            "etiqueta": f"{int(lm) if lm == int(lm) else lm:g} MT",
        }
    return base


def _unidad_compra_label(tipo: str, formato: Dict[str, Any]) -> str:
    """Devuelve el texto de unidad para la columna Unidad de Excel (ej. 'Tramos 12 MT')."""
    if tipo == "perfil":
        lm = float(formato.get("longitud_m") or 6)
        lm_str = str(int(lm)) if lm == int(lm) else f"{lm:g}"
        return f"Tramos {lm_str} MT"
    lp = formato.get("largo_pies")
    ap = formato.get("ancho_pies")
    if lp and ap and float(lp) > 0 and float(ap) > 0:
        lp_s = str(int(float(lp))) if float(lp) == int(float(lp)) else f"{float(lp):g}"
        ap_s = str(int(float(ap))) if float(ap) == int(float(ap)) else f"{float(ap):g}"
        return f"Placas {ap_s}'x{lp_s}'"
    etiq = (formato.get("etiqueta") or "4'x10'").replace("×", "x")
    return f"Placas {etiq}"


def _aplicar_config_compra(
    material_oficial: str,
    sugerencia_auto: str,
    cfg: Optional[Dict[str, Any]],
    req_area_mm2: float,
    req_long_mm: float,
    scrap: float = _SCRAP_COMPRA_DEFAULT,
) -> tuple:
    """Devuelve (sugerencia_final, compra_habilitada, detalle_medidas).

    detalle incluye siempre 'compra_cantidad' (float) y 'compra_unidad' (str)
    para que el Excel pueda mostrar columnas separadas de cantidad y unidad.
    """
    material_upper = material_oficial.upper()
    tipo = _material_tipo_compra(material_upper)
    detalle: Dict[str, Any] = {}

    # ── Sin config: usa detección automática ─────────────────────────────────
    if not cfg:
        if "directa" in sugerencia_auto.lower():
            detalle.update({"directa": True, "compra_cantidad": 0.0, "compra_unidad": "pz"})
            return sugerencia_auto, True, detalle
        if req_area_mm2 <= 0 and req_long_mm <= 0:
            detalle.update({"compra_cantidad": 0.0, "compra_unidad": "—"})
            return sugerencia_auto, True, detalle
        fmt = (
            _detectar_formato_perfil(material_upper)
            if tipo == "perfil"
            else _detectar_formato_placa(material_upper)
        )
        qty = _cantidad_compra_auto(tipo, fmt, req_area_mm2, req_long_mm, scrap)
        detalle.update({
            "compra_cantidad": round(qty, 2),
            "compra_unidad": _unidad_compra_label(tipo, fmt),
        })
        return sugerencia_auto, True, detalle

    cfg = _normalize_cfg_entry(cfg)
    if cfg.get("habilitado") is False:
        detalle.update({"compra_cantidad": 0.0, "compra_unidad": "—"})
        return "— (no comprar)", False, detalle

    # Tipo override: el usuario puede forzar "placa" o "perfil" independientemente
    # del nombre del material, dando control total sobre el cálculo.
    tipo_override = cfg.get("tipo_compra") or ""
    if tipo_override in ("placa", "perfil"):
        tipo = tipo_override

    fmt_id = cfg.get("formato_id") or "auto"

    if fmt_id == "directa":
        texto = (cfg.get("texto") or "").strip()
        detalle.update({"directa": True, "compra_cantidad": 0.0, "compra_unidad": "pz"})
        return texto or "Compra directa por cantidad", True, detalle

    if fmt_id == "manual":
        texto = (cfg.get("texto") or "").strip()
        detalle.update({"compra_cantidad": 0.0, "compra_unidad": ""})
        return texto or sugerencia_auto, True, detalle

    formato = _formato_desde_cfg(cfg, tipo, material_upper)
    if tipo == "placa":
        lp = formato.get("largo_pies")
        ap = formato.get("ancho_pies")
        if lp and ap:
            detalle.update({
                "largo_pies": lp,
                "ancho_pies": ap,
                "area_placa_m2": formato.get("area_m2"),
            })
    else:
        dm = formato.get("longitud_m") or formato.get("distancia_metros")
        if dm:
            detalle["distancia_metros"] = dm

    qty = _cantidad_compra_auto(tipo, formato, req_area_mm2, req_long_mm, scrap)
    texto = _texto_sugerencia_compra(tipo, formato, qty)
    detalle.update({
        "compra_cantidad": round(qty, 2),
        "compra_unidad": _unidad_compra_label(tipo, formato),
    })
    return texto, True, detalle


def _load_all_compra_config() -> Dict[str, Any]:
    if not os.path.isfile(_MRP_COMPRA_CONFIG_PATH):
        return {"global": {}, "revisions": {}}
    try:
        with open(_MRP_COMPRA_CONFIG_PATH, "r", encoding="utf-8") as f:
            data = json.load(f)
        if isinstance(data, dict):
            data.setdefault("global", {})
            data.setdefault("revisions", {})
            return data
    except Exception:
        pass
    return {"global": {}, "revisions": {}}


def _load_material_defaults() -> Dict[str, Dict[str, Any]]:
    """Carga mapeos por defecto (Material Oficial → config de compra) desde JSON."""
    if not os.path.isfile(_MRP_MATERIAL_DEFAULTS_PATH):
        return {}
    try:
        with open(_MRP_MATERIAL_DEFAULTS_PATH, "r", encoding="utf-8") as f:
            raw = json.load(f)
        if not isinstance(raw, dict):
            return {}
        return {
            _material_config_key(str(k)): _normalize_cfg_entry(v)
            for k, v in raw.items()
            if isinstance(v, dict)
        }
    except Exception:
        return {}


def _load_global_compra_config() -> Dict[str, Dict[str, Any]]:
    """Configuración persistente por nombre de material oficial (todas las revisiones).

    Prioridad: config guardada por el usuario > defaults del codebase.
    """
    defaults = _load_material_defaults()
    data = _load_all_compra_config()
    raw = data.get("global", {})
    saved: Dict[str, Dict[str, Any]] = {}
    if isinstance(raw, dict):
        for k, v in raw.items():
            if isinstance(v, dict):
                saved[_material_config_key(str(k))] = _normalize_cfg_entry(v)
    # Saved overrides defaults
    return {**defaults, **saved}


def _save_global_compra_config(config: Dict[str, Dict[str, Any]]) -> None:
    os.makedirs(os.path.dirname(_MRP_COMPRA_CONFIG_PATH), exist_ok=True)
    data = _load_all_compra_config()
    normalized: Dict[str, Dict[str, Any]] = {}
    for k, v in config.items():
        if isinstance(v, dict):
            normalized[_material_config_key(str(k))] = _normalize_cfg_entry(v)
    data["global"] = normalized
    with open(_MRP_COMPRA_CONFIG_PATH, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


def _load_revision_compra_config(id_revision: int) -> Dict[str, Dict[str, Any]]:
    """Compatibilidad: la config efectiva es global por material oficial."""
    return _load_global_compra_config()


def _material_config_key(material_oficial: str) -> str:
    return material_oficial.strip().upper()


class MrpCompraConfigSavePayload(BaseModel):
    # revision_id se conserva por compatibilidad; la config es global (revision_id ignorado).
    revision_id: int = Field(default=0, ge=0)
    config: Dict[str, Dict[str, Any]] = Field(default_factory=dict)


class MrpOptimizarUsoMaterialPayload(BaseModel):
    material_oficial: str = Field(..., min_length=1)
    calibre_espesor: Optional[str] = None
    cantidad_disponible: float = Field(..., gt=0)
    id_revisiones: List[int] = Field(..., min_items=1)
    excluir_si_stock_gt_cero: bool = False
    # Retazo/chapa: medidas en mm. Si ambos > 0, cantidad_disponible = nº de piezas de ese tamaño
    # y el área útil es cantidad * largo * ancho; solo se sugieren piezas que caben (con giro).
    largo_materia_prima_mm: Optional[float] = None
    ancho_materia_prima_mm: Optional[float] = None


def _pieza_cabe_en_placa(
    largo_pieza_mm: float,
    ancho_pieza_mm: float,
    largo_placa_mm: float,
    ancho_placa_mm: float,
) -> bool:
    """Comprueba si la silueta de la pieza cabe en la placa/retazo (eje alineado, rotación 90°)."""
    lp = float(largo_pieza_mm or 0)
    wp = float(ancho_pieza_mm or 0)
    lp_stock = float(largo_placa_mm or 0)
    wp_stock = float(ancho_placa_mm or 0)
    if lp_stock <= 0 or wp_stock <= 0:
        return False
    # Perfil/barra: solo largo útil
    if lp > 0 and wp <= 0:
        return lp <= max(lp_stock, wp_stock)
    if wp > 0 and lp <= 0:
        return wp <= max(lp_stock, wp_stock)
    if lp <= 0 or wp <= 0:
        return False
    return (lp <= lp_stock and wp <= wp_stock) or (lp <= wp_stock and wp <= lp_stock)


def _colocar_rectangulo_en_huecos(
    huecos: List[tuple[float, float]],
    largo_pieza: float,
    ancho_pieza: float,
) -> bool:
    """Coloca una pieza en la lista de huecos con split simple (guillotina) para nesting rápido."""
    mejor_idx = -1
    mejor_orient: Optional[tuple[float, float]] = None
    mejor_desperdicio: Optional[float] = None
    for i, (hw, hh) in enumerate(huecos):
        for pw, ph in ((largo_pieza, ancho_pieza), (ancho_pieza, largo_pieza)):
            if pw <= 0 or ph <= 0:
                continue
            if pw <= hw and ph <= hh:
                desperdicio = (hw * hh) - (pw * ph)
                if mejor_desperdicio is None or desperdicio < mejor_desperdicio:
                    mejor_desperdicio = desperdicio
                    mejor_idx = i
                    mejor_orient = (pw, ph)
    if mejor_idx < 0 or mejor_orient is None:
        return False
    hw, hh = huecos.pop(mejor_idx)
    pw, ph = mejor_orient
    # Split guillotina: derecha + inferior. Es rápido y estable para UI en tiempo real.
    right_w = hw - pw
    bottom_h = hh - ph
    if right_w > 0:
        huecos.append((right_w, ph))
    if bottom_h > 0:
        huecos.append((hw, bottom_h))
    huecos.sort(key=lambda r: r[0] * r[1], reverse=True)
    return True


def _simular_nesting_ligero(
    candidatos: List[Dict[str, Any]],
    largo_placa_mm: float,
    ancho_placa_mm: float,
    cantidad_placas: int,
    warnings: List[str],
) -> Dict[str, Any]:
    """
    Heurística rápida tipo nesting:
    - Simulación por placa con huecos rectangulares (guillotina).
    - Prioriza piezas de mayor área para mejorar aprovechamiento.
    - Limita iteraciones para proteger tiempo de respuesta.
    """
    pendientes: Dict[str, int] = {}
    por_codigo: Dict[str, Dict[str, Any]] = {}
    for c in candidatos:
        cod = str(c.get("codigo_pieza") or "").strip()
        if not cod:
            continue
        por_codigo[cod] = c
        pendientes[cod] = max(0, int(math.floor(float(c.get("demanda_neta") or 0))))
    orden_codigos = sorted(
        por_codigo.keys(),
        key=lambda k: float(por_codigo[k].get("area_unidad_nesting_mm2") or 0),
        reverse=True,
    )
    colocadas: Dict[str, int] = {k: 0 for k in orden_codigos}
    max_colocaciones = 8000
    pasos = 0
    for _ in range(max(0, int(cantidad_placas))):
        huecos: List[tuple[float, float]] = [(float(largo_placa_mm), float(ancho_placa_mm))]
        progreso = True
        while progreso and pasos < max_colocaciones:
            progreso = False
            for cod in orden_codigos:
                if pendientes.get(cod, 0) <= 0:
                    continue
                c = por_codigo[cod]
                l = float(c.get("largo_unidad_mm") or 0)
                a = float(c.get("ancho_unidad_mm") or 0)
                if l <= 0 or a <= 0:
                    continue
                if _colocar_rectangulo_en_huecos(huecos, l, a):
                    pendientes[cod] = max(0, pendientes[cod] - 1)
                    colocadas[cod] = colocadas.get(cod, 0) + 1
                    pasos += 1
                    progreso = True
                if pasos >= max_colocaciones:
                    break
    if pasos >= max_colocaciones:
        warnings.append("Nesting parcial: se alcanzó el límite de iteraciones para mantener respuesta rápida.")
    consumido_mm2 = 0.0
    for cod, qty in colocadas.items():
        if qty <= 0:
            continue
        area_u = float(por_codigo[cod].get("area_unidad_nesting_mm2") or 0)
        consumido_mm2 += qty * area_u
    return {"colocadas": colocadas, "consumido_mm2": consumido_mm2}

@router.get("/api/mrp/revisiones")
def get_mrp_revisiones():
    """Lista DISTINCT de revisiones para el selector del MRPII.
    Usa subconsulta con STRING_AGG para obtener los clientes afectados
    sin multiplicar filas por el JOIN a Tbl_Clientes_Configuracion."""
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("""
            SELECT
                R.ID_Revision,
                TR.Nombre_Tracto,
                TP.Nombre_Tipo,
                V.Nombre_Version,
                R.Numero_Revision,
                ISNULL(R.Estado, '') AS Estado,
                ISNULL(
                    (SELECT STRING_AGG(C.Nombre_Cliente, ', ')
                     FROM Tbl_Clientes_Configuracion C
                     WHERE C.ID_Version = V.ID_Version),
                    'Ingeniería Base (Sin clientes)'
                ) AS Clientes_Afectados
            FROM Tbl_BOM_Revisiones R
            JOIN Tbl_Versiones_Ingenieria V  ON R.ID_Version = V.ID_Version
            JOIN Tbl_Tipos_Proyecto      TP  ON V.ID_Tipo    = TP.ID_Tipo
            JOIN Tbl_Proyectos_Tracto    TR  ON TP.ID_Tracto = TR.ID_Tracto
            -- Solo revisiones Aprobadas en el selector MRPII.
            -- Las OBSOLETAS y Borradores no deben usarse para cálculo de materiales.
            WHERE R.Estado = 'Aprobada'
            ORDER BY TR.Nombre_Tracto, TP.Nombre_Tipo, V.Nombre_Version, R.Numero_Revision
        """)
        return [
            {
                "id_revision":         int(r.ID_Revision),
                "nombre_tracto":       r.Nombre_Tracto       or "",
                "nombre_tipo":         r.Nombre_Tipo         or "",
                "nombre_version":      r.Nombre_Version      or "",
                "numero_revision":     r.Numero_Revision,
                "estado":              r.Estado              or "",
                "clientes_afectados":  r.Clientes_Afectados  or "Ingeniería Base (Sin clientes)",
            }
            for r in cursor.fetchall()
        ]
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        conn.close()


@router.post("/api/mrp/optimizar_uso_material")
def optimizar_uso_material(payload: MrpOptimizarUsoMaterialPayload):
    revisiones = sorted({int(x) for x in payload.id_revisiones if int(x) > 0})
    if not revisiones:
        raise HTTPException(status_code=400, detail="id_revisiones inválido")
    material = (payload.material_oficial or "").strip()
    if not material:
        raise HTTPException(status_code=400, detail="material_oficial es obligatorio")
    calibre = (payload.calibre_espesor or "").strip()
    disponible = float(payload.cantidad_disponible or 0)
    if disponible <= 0:
        raise HTTPException(status_code=400, detail="cantidad_disponible debe ser mayor a 0")
    lm = float(payload.largo_materia_prima_mm or 0)
    wm = float(payload.ancho_materia_prima_mm or 0)
    if (lm > 0) != (wm > 0):
        raise HTTPException(
            status_code=400,
            detail="Indique largo y ancho de la materia prima (mm), o deje ambos vacíos para usar solo superficie total (mm²).",
        )
    usa_medidas_placa = lm > 0 and wm > 0
    if usa_medidas_placa:
        pool_mm2 = disponible * lm * wm
    else:
        pool_mm2 = disponible
    if pool_mm2 <= 0:
        raise HTTPException(status_code=400, detail="Superficie disponible inválida (revisar cantidad y medidas)")

    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        placeholders = ",".join(["?"] * len(revisiones))
        where_calibre = ""
        params: List[Any] = list(revisiones)
        params.append(material)
        if calibre:
            where_calibre = " AND ISNULL(CAST(M.Espesor_Perfil_CAD AS VARCHAR(50)), 'N/A') = ? "
            params.append(calibre)

        query = f"""
        WITH PiezasBase AS (
            SELECT
                E.Codigo_Pieza,
                ISNULL(NULLIF(LTRIM(RTRIM(M.Descripcion)), ''), E.Codigo_Pieza) AS Descripcion,
                LTRIM(RTRIM(ISNULL(M.Material, ''))) AS Material_Oficial,
                ISNULL(CAST(M.Espesor_Perfil_CAD AS VARCHAR(50)), 'N/A') AS Calibre_Espesor,
                ISNULL(M.Stock_PT_Almacen, 0) AS Stock_PT_Almacen,
                ISNULL(E.Cantidad, 0) AS Cantidad,
                COALESCE(
                    TRY_CAST(M.Largo_CAD AS FLOAT),
                    TRY_CAST(REPLACE(REPLACE(REPLACE(REPLACE(M.Largo_CAD,' mm',''),',',''),' ',''),'-','') AS FLOAT),
                    TRY_CAST(M.Largo_DXF AS FLOAT),
                    TRY_CAST(REPLACE(REPLACE(REPLACE(REPLACE(M.Largo_DXF,' mm',''),',',''),' ',''),'-','') AS FLOAT)
                ) AS LargoLimpio,
                COALESCE(
                    TRY_CAST(M.Ancho_CAD AS FLOAT),
                    TRY_CAST(REPLACE(REPLACE(REPLACE(REPLACE(M.Ancho_CAD,' mm',''),',',''),' ',''),'-','') AS FLOAT),
                    TRY_CAST(M.Ancho_DXF AS FLOAT),
                    TRY_CAST(REPLACE(REPLACE(REPLACE(REPLACE(M.Ancho_DXF,' mm',''),',',''),' ',''),'-','') AS FLOAT)
                ) AS AnchoLimpio,
                COALESCE(
                    TRY_CAST(M.Area_CAD AS FLOAT),
                    TRY_CAST(REPLACE(REPLACE(REPLACE(REPLACE(M.Area_CAD,' mm^2',''),',',''),' ',''),'-','') AS FLOAT)
                ) AS AreaLimpia
            FROM Tbl_BOM_Estructura E
            JOIN Tbl_Ensambles EN ON E.ID_Ensamble = EN.ID_Ensamble
            JOIN Tbl_Estaciones ES ON EN.ID_Estacion = ES.ID_Estacion
            LEFT JOIN Tbl_Maestro_Piezas M
                ON LTRIM(RTRIM(E.Codigo_Pieza)) = LTRIM(RTRIM(M.Codigo_Pieza))
            WHERE ES.ID_Revision IN ({placeholders})
              AND UPPER(LTRIM(RTRIM(ISNULL(M.Material, '')))) NOT LIKE '%COMERCIAL%'
              AND UPPER(LTRIM(RTRIM(ISNULL(M.Medida, '')))) <> 'COMERCIAL'
              AND LTRIM(RTRIM(ISNULL(M.Material, ''))) = ?
              {where_calibre}
        )
        SELECT
            Codigo_Pieza,
            MAX(Descripcion) AS Descripcion,
            MAX(Material_Oficial) AS Material_Oficial,
            MAX(Calibre_Espesor) AS Calibre_Espesor,
            SUM(Cantidad) AS Demanda_Total,
            MAX(Stock_PT_Almacen) AS Stock_PT_Almacen,
            SUM(Cantidad * ISNULL(NULLIF(AreaLimpia, 0), (ISNULL(LargoLimpio, 0) * ISNULL(AnchoLimpio, 0)))) AS Requerimiento_Area_mm2,
            SUM(Cantidad * ISNULL(LargoLimpio, 0)) AS Requerimiento_Longitud_mm,
            MAX(ISNULL(LargoLimpio, 0)) AS Largo_Unidad_mm,
            MAX(ISNULL(AnchoLimpio, 0)) AS Ancho_Unidad_mm
        FROM PiezasBase
        GROUP BY Codigo_Pieza
        ORDER BY Codigo_Pieza
        """
        cursor.execute(query, tuple(params))
        rows = cursor.fetchall()

        overstock: List[Dict[str, Any]] = []
        candidatos: List[Dict[str, Any]] = []
        descartadas_geometria: List[Dict[str, Any]] = []
        warnings: List[str] = []
        for r in rows:
            codigo = str(r.Codigo_Pieza or "").strip()
            demanda = float(r.Demanda_Total or 0)
            stock = float(r.Stock_PT_Almacen or 0)
            area = float(r.Requerimiento_Area_mm2 or 0)
            long_mm = float(r.Requerimiento_Longitud_mm or 0)
            consumo_unit = (area / demanda) if (demanda > 0 and area > 0) else ((long_mm / demanda) if (demanda > 0 and long_mm > 0) else 0.0)
            if consumo_unit <= 0:
                warnings.append(f"{codigo}: sin medidas CAD/DXF; no se optimiza")
                continue
            excluida = stock > 0 if payload.excluir_si_stock_gt_cero else stock >= demanda
            if excluida:
                overstock.append(
                    {
                        "codigo_pieza": codigo,
                        "descripcion": str(r.Descripcion or codigo),
                        "stock_pt": stock,
                        "demanda_en_revisiones": demanda,
                        "motivo": "sobrestock",
                    }
                )
                continue
            demanda_neta = max(0.0, demanda - stock)
            if demanda_neta <= 0:
                continue
            largo_u = float(getattr(r, "Largo_Unidad_mm", None) or 0)
            ancho_u = float(getattr(r, "Ancho_Unidad_mm", None) or 0)
            area_u_nesting = (largo_u * ancho_u) if (largo_u > 0 and ancho_u > 0) else 0.0
            if area_u_nesting <= 0 and demanda > 0 and area > 0:
                area_u_nesting = area / demanda
            if usa_medidas_placa:
                if largo_u <= 0 and ancho_u <= 0:
                    warnings.append(f"{codigo}: sin largo/ancho en maestro; no se puede verificar corte en placa")
                    descartadas_geometria.append(
                        {
                            "codigo_pieza": codigo,
                            "descripcion": str(r.Descripcion or codigo),
                            "motivo": "sin_medidas_cad",
                        }
                    )
                    continue
                if not _pieza_cabe_en_placa(largo_u, ancho_u, lm, wm):
                    descartadas_geometria.append(
                        {
                            "codigo_pieza": codigo,
                            "descripcion": str(r.Descripcion or codigo),
                            "motivo": "no_cabe_en_placa",
                            "largo_pieza_mm": largo_u,
                            "ancho_pieza_mm": ancho_u,
                        }
                    )
                    continue
            candidatos.append(
                {
                    "codigo_pieza": codigo,
                    "descripcion": str(r.Descripcion or codigo),
                    "calibre_espesor": str(r.Calibre_Espesor or "N/A"),
                    "demanda_neta": demanda_neta,
                    "consumo_unitario": consumo_unit,
                    "stock_pt": stock,
                    "material": str(r.Material_Oficial or material),
                    "requerimiento_material_por_unidad": consumo_unit,
                    "largo_unidad_mm": largo_u,
                    "ancho_unidad_mm": ancho_u,
                    "area_unidad_nesting_mm2": area_u_nesting,
                }
            )

        candidatos.sort(key=lambda x: (x["consumo_unitario"], x["codigo_pieza"]))
        fabricables: List[Dict[str, Any]] = []
        restante = pool_mm2
        if usa_medidas_placa:
            placas_enteras = int(math.floor(disponible))
            if placas_enteras <= 0:
                raise HTTPException(
                    status_code=400,
                    detail="Con largo/ancho de placa, cantidad_disponible debe ser al menos 1.",
                )
            if not math.isclose(disponible, float(placas_enteras)):
                warnings.append(
                    "cantidad_disponible con decimales: nesting usa solo placas completas para preservar rendimiento."
                )
            sim = _simular_nesting_ligero(candidatos, lm, wm, placas_enteras, warnings)
            colocadas = sim.get("colocadas", {})
            consumido_total = float(sim.get("consumido_mm2") or 0)
            restante = max(0.0, pool_mm2 - consumido_total)
            for c in candidatos:
                codigo = str(c["codigo_pieza"])
                sugerida = int(colocadas.get(codigo, 0) or 0)
                if sugerida <= 0:
                    continue
                consumo_unit = float(c.get("area_unidad_nesting_mm2") or c["consumo_unitario"] or 0)
                consumido = sugerida * consumo_unit
                fabricables.append(
                    {
                        "codigo_pieza": c["codigo_pieza"],
                        "descripcion": c["descripcion"],
                        "cantidad_bom_agregada": c["demanda_neta"],
                        "requerimiento_material_por_unidad": consumo_unit,
                        "cantidad_sugerida_fabricar": sugerida,
                        "material_consumido_estimado": consumido,
                        "stock_pt": c["stock_pt"],
                    }
                )
        else:
            for c in candidatos:
                max_units = math.floor(restante / c["consumo_unitario"]) if c["consumo_unitario"] > 0 else 0
                sugerida = min(int(c["demanda_neta"]), int(max_units))
                if sugerida <= 0:
                    continue
                consumido = sugerida * c["consumo_unitario"]
                restante = max(0.0, restante - consumido)
                fabricables.append(
                    {
                        "codigo_pieza": c["codigo_pieza"],
                        "descripcion": c["descripcion"],
                        "cantidad_bom_agregada": c["demanda_neta"],
                        "requerimiento_material_por_unidad": c["requerimiento_material_por_unidad"],
                        "cantidad_sugerida_fabricar": sugerida,
                        "material_consumido_estimado": consumido,
                        "stock_pt": c["stock_pt"],
                    }
                )

        return {
            "material_oficial": material,
            "calibre_espesor": calibre or None,
            "cantidad_disponible": disponible,
            "largo_materia_prima_mm": lm if usa_medidas_placa else None,
            "ancho_materia_prima_mm": wm if usa_medidas_placa else None,
            "area_disponible_mm2": pool_mm2,
            "fabricables": fabricables,
            "sobrestock": overstock,
            "descartadas_por_medida": descartadas_geometria,
            "resumen": {
                "material_restante": restante,
                "material_consumido": max(0.0, pool_mm2 - restante),
                "piezas_consideradas": len(candidatos),
                "piezas_fabricables": len(fabricables),
                "warnings": warnings,
                "modo_medidas_placa": usa_medidas_placa,
            },
        }
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        conn.close()


@router.get("/api/mrp/compra-formatos")
def get_compra_formatos():
    return {"formatos": FORMATOS_COMPRA_MP}


@router.get("/api/mrp/compra-config/{id_revision}")
def get_compra_config(id_revision: int):
    return {
        "revision_id": id_revision,
        "config": _load_global_compra_config(),
        "formatos": FORMATOS_COMPRA_MP,
        "scope": "global",
    }


@router.put("/api/mrp/compra-config")
def save_compra_config(payload: MrpCompraConfigSavePayload):
    _save_global_compra_config(payload.config or {})
    cfg = _load_global_compra_config()
    return {"ok": True, "scope": "global", "count": len(cfg)}


@router.get("/api/mrp/calculate/{id_revision}")
def calculate_mrp(id_revision: int):
    """Calcula la consolidación de compras (MRP) con filtrado estricto y diagnóstico de huérfanos.
    Separa los Componentes Comerciales del cálculo de placas/perfiles."""
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        # ── CTE base compartida (reutilizada en las tres queries) ─────────────
        # PURGA REGLA ESPEJO: se eliminó M.Descripcion completamente.
        # Material se extrae ÚNICA y EXCLUSIVAMENTE de M.Material.
        # JOIN es LEFT JOIN para no perder piezas que aún no tienen CAD.
        _cte_base = """
        WITH PiezasBase AS (
            SELECT
                E.Codigo_Pieza,
                NULLIF(LTRIM(RTRIM(M.Material)), '')              AS MaterialOficialRaw,
                LTRIM(RTRIM(ISNULL(M.Material, '')))              AS Material_Trace,
                LTRIM(RTRIM(ISNULL(M.Medida, '')))                AS Medida_Trace,
                CASE
                    WHEN UPPER(LTRIM(RTRIM(ISNULL(M.Material, '')))) LIKE '%COMERCIAL%'
                      OR UPPER(LTRIM(RTRIM(ISNULL(M.Medida, '')))) = 'COMERCIAL'
                    THEN 1 ELSE 0
                END                                               AS EsComercial,
                M.Espesor_Perfil_CAD,
                ISNULL(M.Stock_PT_Almacen, 0)                     AS Stock_PT_Almacen,
                M.Stock_PT_Almacen_SyncAt                         AS Stock_PT_Almacen_SyncAt,
                ISNULL(E.Cantidad, 0)                             AS Cantidad,
                COALESCE(
                    TRY_CAST(M.Largo_CAD AS FLOAT),
                    TRY_CAST(REPLACE(REPLACE(REPLACE(REPLACE(M.Largo_CAD,' mm',''),',',''),' ',''),'-','') AS FLOAT),
                    TRY_CAST(M.Largo_DXF AS FLOAT),
                    TRY_CAST(REPLACE(REPLACE(REPLACE(REPLACE(M.Largo_DXF,' mm',''),',',''),' ',''),'-','') AS FLOAT)
                ) AS LargoLimpio,
                COALESCE(
                    TRY_CAST(M.Ancho_CAD AS FLOAT),
                    TRY_CAST(REPLACE(REPLACE(REPLACE(REPLACE(M.Ancho_CAD,' mm',''),',',''),' ',''),'-','') AS FLOAT),
                    TRY_CAST(M.Ancho_DXF AS FLOAT),
                    TRY_CAST(REPLACE(REPLACE(REPLACE(REPLACE(M.Ancho_DXF,' mm',''),',',''),' ',''),'-','') AS FLOAT)
                ) AS AnchoLimpio,
                COALESCE(
                    TRY_CAST(M.Area_CAD AS FLOAT),
                    TRY_CAST(REPLACE(REPLACE(REPLACE(REPLACE(M.Area_CAD,' mm^2',''),',',''),' ',''),'-','') AS FLOAT)
                ) AS AreaLimpia
            FROM Tbl_BOM_Estructura E
            JOIN Tbl_Ensambles   EN ON E.ID_Ensamble  = EN.ID_Ensamble
            JOIN Tbl_Estaciones  ES ON EN.ID_Estacion = ES.ID_Estacion
            LEFT JOIN Tbl_Maestro_Piezas M
                ON LTRIM(RTRIM(E.Codigo_Pieza)) = LTRIM(RTRIM(M.Codigo_Pieza))
            WHERE ES.ID_Revision = ?
        )
        """

        # ── 1. Materia Prima / Placas (excluye COMERCIAL) ────────────────────
        # FIX MRP-ESPESOR-001: se elimina Espesor_Perfil_CAD del GROUP BY.
        # La agrupación es ahora por (material_oficial, Codigo_Pieza).
        # La consolidación final por material se hace en Python, donde se
        # deriva el calibre canónico del NOMBRE del material (_calibre_canonico),
        # evitando que variaciones de medición CAD fragmenten la orden de compra.
        #
        # FIX ORPHAN-EXCLUSION: piezas sin dimensiones CAD (LargoLimpio = 0 Y
        # AreaLimpia = 0) se excluyen de BaseNoCommercial.  Dichas piezas ya
        # aparecen en query_orphans y sólo deben mostrarse en la pestaña
        # "Auditoría / Huérfanos" hasta que se les asignen dimensiones.
        _dc_in = _directa_codigos_sql_in()
        query_mrp = _cte_base + f"""
        , BaseNoCommercial AS (
            SELECT *
            FROM PiezasBase
            WHERE MaterialOficialRaw IS NOT NULL
              AND EsComercial = 0
              AND (ISNULL(LargoLimpio, 0) > 0 OR ISNULL(AreaLimpia, 0) > 0)
              AND Codigo_Pieza NOT IN {_dc_in}
        ),
        AggPerCode AS (
            SELECT
                LTRIM(RTRIM(MaterialOficialRaw))                            AS material_oficial,
                Codigo_Pieza,
                SUM(Cantidad)                                               AS Cantidad_Piezas,
                SUM(Cantidad * ISNULL(LargoLimpio, 0.0))                    AS Requerimiento_Longitud_mm,
                SUM(Cantidad * ISNULL(NULLIF(AreaLimpia, 0),
                    (ISNULL(LargoLimpio, 0) * ISNULL(AnchoLimpio, 0))))     AS Requerimiento_Area_mm2,
                MAX(ISNULL(Stock_PT_Almacen, 0))                            AS Stock_PT_Almacen,
                MAX(Stock_PT_Almacen_SyncAt)                                AS Stock_PT_Almacen_SyncAt,
                MAX(LargoLimpio)                                             AS Largo_Pieza_mm,
                MAX(AnchoLimpio)                                             AS Ancho_Pieza_mm
            FROM BaseNoCommercial
            GROUP BY
                LTRIM(RTRIM(MaterialOficialRaw)),
                Codigo_Pieza
        )
        SELECT
            material_oficial,
            Codigo_Pieza,
            Cantidad_Piezas,
            Requerimiento_Longitud_mm,
            Requerimiento_Area_mm2,
            Stock_PT_Almacen,
            Stock_PT_Almacen_SyncAt,
            Largo_Pieza_mm,
            Ancho_Pieza_mm
        FROM AggPerCode
        ORDER BY material_oficial, Codigo_Pieza
        """
        cursor.execute(query_mrp, (id_revision,))
        rows_mrp = cursor.fetchall()

        # ── Consolidación en Python por material_oficial ───────────────────
        # Agrupamos todas las piezas bajo su material, sumando totales y
        # construyendo la lista de piezas hija para la UI expandable.
        from collections import OrderedDict as _OD

        _piezas_por_mat: Dict[str, List[Dict[str, Any]]] = _OD()
        _area_por_mat: Dict[str, float] = {}
        _long_por_mat: Dict[str, float] = {}
        _cant_por_mat: Dict[str, float] = {}
        _stock_por_mat: Dict[str, float] = {}
        _sync_por_mat: Dict[str, Any] = {}

        for r in rows_mrp:
            mat_of = str(r.material_oficial or "").strip()
            if mat_of not in _piezas_por_mat:
                _piezas_por_mat[mat_of] = []
                _area_por_mat[mat_of]   = 0.0
                _long_por_mat[mat_of]   = 0.0
                _cant_por_mat[mat_of]   = 0.0
                _stock_por_mat[mat_of]  = 0.0

            cant   = float(r.Cantidad_Piezas or 0)
            area   = float(r.Requerimiento_Area_mm2 or 0)
            long   = float(r.Requerimiento_Longitud_mm or 0)
            stock  = float(r.Stock_PT_Almacen or 0)
            sync   = r.Stock_PT_Almacen_SyncAt
            largo  = float(getattr(r, "Largo_Pieza_mm", None) or 0)
            ancho  = float(getattr(r, "Ancho_Pieza_mm", None) or 0)
            area_u = area / cant if cant > 0 else 0.0

            _piezas_por_mat[mat_of].append({
                "codigo_pieza":      str(r.Codigo_Pieza or "").strip(),
                "cantidad":          cant,
                "area_mm2":          area,
                "area_unitaria_mm2": area_u,
                "longitud_mm":       long,
                "largo_mm":          largo,
                "ancho_mm":          ancho,
                "stock_pieza":       stock,
            })
            _area_por_mat[mat_of]  += area
            _long_por_mat[mat_of]  += long
            _cant_por_mat[mat_of]  += cant
            _stock_por_mat[mat_of] += stock
            if sync:
                prev = _sync_por_mat.get(mat_of)
                if prev is None or str(sync) > str(prev):
                    _sync_por_mat[mat_of] = sync

        mrp_calculado = []
        compra_cfg_rev = _load_global_compra_config()
        scrap_factor = _SCRAP_COMPRA_DEFAULT

        for mat_of, piezas_hijas in _piezas_por_mat.items():
            calibre      = _calibre_canonico(mat_of)
            req_area_mm2 = _area_por_mat[mat_of]
            req_long_mm  = _long_por_mat[mat_of]
            total_piezas = _cant_por_mat[mat_of]
            total_stock  = _stock_por_mat[mat_of]
            brecha_estim = max(0.0, total_piezas - total_stock)
            sync_at      = _sync_por_mat.get(mat_of)

            sugerencia_auto = _sugerencia_compra_automatica(
                mat_of, req_area_mm2, req_long_mm, scrap_factor
            )
            cfg_mat = compra_cfg_rev.get(_material_config_key(mat_of))
            sugerencia, compra_ok, detalle_compra = _aplicar_config_compra(
                mat_of,
                sugerencia_auto,
                cfg_mat,
                req_area_mm2,
                req_long_mm,
                scrap_factor,
            )

            # Para compra directa la cantidad es la brecha BOM, no área calculada
            compra_cantidad: float = detalle_compra.get("compra_cantidad") or 0.0
            compra_unidad: str = detalle_compra.get("compra_unidad") or ""
            if detalle_compra.get("directa"):
                compra_cantidad = float(brecha_estim)

            mrp_calculado.append({
                "material_oficial":          mat_of,
                "Material":                  mat_of,
                "Calibre_Espesor":           calibre,
                "Cantidad_Total_Piezas":     total_piezas,
                "Requerimiento_Area_mm2":    req_area_mm2,
                "Requerimiento_Longitud_mm": req_long_mm,
                "Sugerencia_Compra":         sugerencia,
                "Sugerencia_Compra_Auto":    sugerencia_auto,
                "Compra_Habilitada":         compra_ok,
                # Si el usuario configuró tipo_compra manualmente, reportarlo; si no, auto.
                "Tipo_Compra":               (
                    (cfg_mat or {}).get("tipo_compra")
                    or _material_tipo_compra(mat_of.upper())
                ),
                "Compra_Largo_Pies":         detalle_compra.get("largo_pies"),
                "Compra_Ancho_Pies":         detalle_compra.get("ancho_pies"),
                "Compra_Area_Placa_m2":      detalle_compra.get("area_placa_m2"),
                "Compra_Distancia_Metros":   detalle_compra.get("distancia_metros"),
                "Compra_Cantidad":           compra_cantidad,
                "Compra_Unidad":             compra_unidad,
                "Stock_Asociado_Estimado":   total_stock,
                "Brecha_Estimada":           brecha_estim,
                "es_estimado":               True,
                "Stock_PT_Almacen_SyncAt":   sync_at.isoformat() if sync_at else None,
                "piezas":                    piezas_hijas,
            })

        # ── 1b. Piezas de compra directa por código (JE-012, etc.) ─────────────
        # Estas piezas no tienen cálculo de área/perfil; se compran por cantidad BOM.
        if _DIRECTA_CODIGOS:
            query_directa_cod = _cte_base + f"""
            SELECT
                Codigo_Pieza,
                SUM(Cantidad) AS Cantidad_Total,
                MAX(ISNULL(Stock_PT_Almacen, 0)) AS Stock_PT_Almacen,
                MAX(Stock_PT_Almacen_SyncAt)     AS Stock_PT_Almacen_SyncAt
            FROM PiezasBase
            WHERE Codigo_Pieza IN {_dc_in}
            GROUP BY Codigo_Pieza
            """
            cursor.execute(query_directa_cod, (id_revision,))
            rows_dc = cursor.fetchall()

            # Obtener descripción desde Maestro
            _ph_dc = ", ".join("?" for _ in _DIRECTA_CODIGOS)
            cursor.execute(
                f"SELECT LTRIM(RTRIM(Codigo_Pieza)), "
                f"LTRIM(RTRIM(ISNULL(Descripcion, Codigo_Pieza))) "
                f"FROM Tbl_Maestro_Piezas "
                f"WHERE LTRIM(RTRIM(Codigo_Pieza)) IN ({_ph_dc})",
                sorted(_DIRECTA_CODIGOS),
            )
            _dc_desc: Dict[str, str] = {r[0]: r[1] for r in cursor.fetchall()}

            for r in rows_dc:
                codigo   = str(r.Codigo_Pieza or "").strip()
                desc     = _dc_desc.get(codigo, codigo)
                cantidad = float(r.Cantidad_Total or 0)
                stock    = float(getattr(r, "Stock_PT_Almacen", 0) or 0)
                brecha   = max(0.0, cantidad - stock)
                sync_at  = getattr(r, "Stock_PT_Almacen_SyncAt", None)
                mrp_calculado.append({
                    "material_oficial":          desc or codigo,
                    "Material":                  desc or codigo,
                    # Almacenar el código en Calibre_Espesor para mostrarlo en Excel
                    "Calibre_Espesor":           codigo,
                    "Cantidad_Total_Piezas":     cantidad,
                    "Requerimiento_Area_mm2":    0.0,
                    "Requerimiento_Longitud_mm": 0.0,
                    "Sugerencia_Compra":         "Compra directa por cantidad",
                    "Sugerencia_Compra_Auto":    "Compra directa por cantidad",
                    "Compra_Habilitada":         True,
                    "Tipo_Compra":               "directa",
                    "Compra_Largo_Pies":         None,
                    "Compra_Ancho_Pies":         None,
                    "Compra_Area_Placa_m2":      None,
                    "Compra_Distancia_Metros":   None,
                    "Compra_Cantidad":           brecha,
                    "Compra_Unidad":             "pz",
                    "Stock_Asociado_Estimado":   stock,
                    "Brecha_Estimada":           brecha,
                    "es_estimado":               False,
                    "Stock_PT_Almacen_SyncAt":   sync_at.isoformat() if sync_at else None,
                    "piezas":                    [],
                })

        # ── 2. Componentes Comerciales (solo cantidad, sin placas) ─────────────
        # Catálogo: Medida = 'COMERCIAL' y/o Material contiene 'COMERCIAL'.
        query_comerciales = _cte_base + f"""
        SELECT
            Codigo_Pieza,
            LTRIM(RTRIM(ISNULL(Material_Trace, '')))  AS Material_Comercial,
            SUM(Cantidad)                              AS Cantidad_Total,
            MAX(ISNULL(Stock_PT_Almacen, 0))          AS Stock_PT_Almacen,
            MAX(Stock_PT_Almacen_SyncAt)              AS Stock_PT_Almacen_SyncAt
        FROM PiezasBase
        WHERE EsComercial = 1
          AND Codigo_Pieza NOT IN {_dc_in}
        GROUP BY Codigo_Pieza, LTRIM(RTRIM(ISNULL(Material_Trace, '')))
        ORDER BY Material_Comercial, Codigo_Pieza
        """
        cursor.execute(query_comerciales, (id_revision,))
        rows_com = cursor.fetchall()

        componentes_comerciales = []
        for r in rows_com:
            demanda = float(r.Cantidad_Total or 0)
            stock = float(getattr(r, "Stock_PT_Almacen", 0) or 0)
            faltante = max(0.0, demanda - stock)
            cobertura_pct = min(100.0, (stock / demanda) * 100.0) if demanda > 0 else 100.0
            sync_at = getattr(r, "Stock_PT_Almacen_SyncAt", None)
            componentes_comerciales.append(
                {
                    "Codigo_Pieza": r.Codigo_Pieza,
                    "Descripcion": r.Material_Comercial,  # campo renombrado; UI compat
                    "Cantidad_Total": demanda,
                    "Stock_PT_Almacen": stock,
                    "Cantidad_Faltante": faltante,
                    "Cobertura_Pct": cobertura_pct,
                    "Stock_PT_Almacen_SyncAt": sync_at.isoformat() if sync_at else None,
                }
            )

        # ── 3. Piezas sin medidas / sin material (huérfanas) ─────────────────
        # PURGA REGLA ESPEJO: huérfano = Material vacío, NULL o 'POR DEFINIR'.
        query_orphans = _cte_base + f"""
        SELECT
            Codigo_Pieza,
            ISNULL((SELECT TOP 1 Nombre_Ensamble
                    FROM Tbl_Ensambles EN2
                    JOIN Tbl_BOM_Estructura E2 ON E2.ID_Ensamble = EN2.ID_Ensamble
                    WHERE E2.Codigo_Pieza = PiezasBase.Codigo_Pieza), 'N/A') AS Nombre_Ensamble,
            ISNULL(Material_Trace, '')  AS Material,
            Cantidad,
            CASE
                WHEN MaterialOficialRaw IS NULL
                     AND (ISNULL(AreaLimpia,0)=0 AND ISNULL(LargoLimpio,0)=0
                          AND (ISNULL(LargoLimpio,0)*ISNULL(AnchoLimpio,0))=0)
                     THEN 'Sin Material ni Dimensiones'
                WHEN MaterialOficialRaw IS NULL THEN 'Falta Asignar Material'
                ELSE 'Sin Dimensiones CAD'
            END AS Motivo_Rechazo
        FROM PiezasBase
        WHERE (MaterialOficialRaw IS NULL
               OR ISNULL(AreaLimpia,0) + ISNULL(LargoLimpio,0) = 0)
          -- Excluir componentes comerciales (Material o Medida = COMERCIAL en maestro).
          AND EsComercial = 0
          -- Excluir piezas de compra directa por código: aparecen en COMPRA DIRECTA.
          AND Codigo_Pieza NOT IN {_dc_in}
        """
        cursor.execute(query_orphans, (id_revision,))
        rows_orphans = cursor.fetchall()

        piezas_sin_medidas = [
            {
                "Codigo_Pieza":   r.Codigo_Pieza,
                "Nombre_Ensamble": r.Nombre_Ensamble,
                "Material":        r.Material,
                "Cantidad":        float(r.Cantidad),
                "Motivo_Rechazo":  r.Motivo_Rechazo,
            }
            for r in rows_orphans
        ]

        # 4. Resumen de stock para widgets inferiores/laterales (MRP)
        demanda_total_com = sum(float(x.get("Cantidad_Total", 0) or 0) for x in componentes_comerciales)
        stock_total_com = sum(float(x.get("Stock_PT_Almacen", 0) or 0) for x in componentes_comerciales)
        faltante_total_com = sum(float(x.get("Cantidad_Faltante", 0) or 0) for x in componentes_comerciales)
        lineas_faltante_com = sum(1 for x in componentes_comerciales if float(x.get("Cantidad_Faltante", 0) or 0) > 0)

        demanda_total_mp = sum(float(x.get("Cantidad_Total_Piezas", 0) or 0) for x in mrp_calculado)
        stock_total_mp = sum(float(x.get("Stock_Asociado_Estimado", 0) or 0) for x in mrp_calculado)
        brecha_total_mp = sum(float(x.get("Brecha_Estimada", 0) or 0) for x in mrp_calculado)
        lineas_brecha_mp = sum(1 for x in mrp_calculado if float(x.get("Brecha_Estimada", 0) or 0) > 0)

        ultima_sync_candidates = [
            x.get("Stock_PT_Almacen_SyncAt")
            for x in componentes_comerciales + mrp_calculado
            if x.get("Stock_PT_Almacen_SyncAt")
        ]
        ultima_sync_stock_pt = max(ultima_sync_candidates) if ultima_sync_candidates else None

        return {
            "mrp_calculado":          mrp_calculado,
            "componentes_comerciales": componentes_comerciales,
            "piezas_sin_medidas":     piezas_sin_medidas,
            "compra_config":          compra_cfg_rev,
            "formatos_compra":        FORMATOS_COMPRA_MP,
            "resumen_stock_mrp": {
                "comerciales": {
                    "demanda_total_unidades": demanda_total_com,
                    "stock_total_unidades": stock_total_com,
                    "faltante_total_unidades": faltante_total_com,
                    "lineas_con_faltante": lineas_faltante_com,
                },
                "materia_prima_estimado": {
                    "demanda_total_unidades": demanda_total_mp,
                    "stock_asociado_total_unidades": stock_total_mp,
                    "brecha_total_unidades": brecha_total_mp,
                    "lineas_con_brecha": lineas_brecha_mp,
                    "es_estimado": True,
                },
                "ultima_sync_stock_pt": ultima_sync_stock_pt,
            },
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        conn.close()

