"""API router: qa."""
import ast
import io
import json
import math
import os
import re
import shutil
import socket
import subprocess
import sys
import traceback
import uuid
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

import openpyxl
import pandas as pd
import pyodbc
from fastapi import APIRouter, BackgroundTasks, File, Form, Header, HTTPException, Request, Response, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from database import get_db_connection, _int_from_count_row
from models import *
from admin_master_password import assert_admin_master_password_matches
from audit_service import registrar_log_global

router = APIRouter()


def _get_cols(cur: Any, table: str) -> List[str]:
    cur.execute(
        """
        SELECT COLUMN_NAME
        FROM INFORMATION_SCHEMA.COLUMNS
        WHERE TABLE_NAME = ?
        ORDER BY ORDINAL_POSITION
        """,
        (table,),
    )
    return [r[0] for r in cur.fetchall()]


def _pick(cols: List[str], *candidates: str) -> Optional[str]:
    lookup = {c.lower(): c for c in cols}
    for cand in candidates:
        hit = lookup.get(cand.lower())
        if hit:
            return hit
    return None


def _tabla_existe(cur: Any, name: str) -> bool:
    cur.execute("SELECT 1 FROM INFORMATION_SCHEMA.TABLES WHERE TABLE_NAME = ?", (name,))
    return cur.fetchone() is not None


def _crear_tarea_correccion_ingenieria(
    cur: Any,
    id_reporte: int,
    payload: BugReportPayload,
) -> Optional[int]:
    t_cols = _get_cols(cur, "Tbl_Gestor_Tareas")
    c_cols = _get_cols(cur, "Tbl_Gestor_Checklist")
    if not t_cols or not c_cols:
        return None
    t_pk = _pick(t_cols, "ID_Tarea", "Id_Tarea")
    c_fk = _pick(c_cols, "ID_Tarea", "Id_Tarea", "Tarea_ID")
    c_nombre = _pick(c_cols, "Nombre_Item", "Item", "Descripcion")
    c_done = _pick(c_cols, "Completado", "Hecho", "Status")
    if not t_pk or not c_fk:
        return None

    cur.execute(
        """
        SELECT Username FROM Tbl_Usuarios
        WHERE UPPER(Rol) IN ('INGENIERIA_METODOS', 'INGENIERIA')
        ORDER BY Username
        """
    )
    usuarios = [str(r[0]).strip() for r in cur.fetchall() if str(r[0] or "").strip()]
    if not usuarios:
        return None
    responsable = usuarios[0]
    titulo = f"[QA] {payload.modulo} · {payload.gravedad}"
    desc = payload.descripcion.strip()
    if payload.contexto_pantalla:
        desc = f"{desc}\n\nContexto: {payload.contexto_pantalla.strip()}"

    col_titulo = _pick(t_cols, "Titulo", "Nombre_Tarea")
    col_desc = _pick(t_cols, "Descripcion", "Detalle")
    col_estado = _pick(t_cols, "Estado", "Status")
    col_prog = _pick(t_cols, "Porcentaje_Progreso", "Progreso")
    col_tipo = _pick(t_cols, "Tipo_Tarea", "Tipo")
    col_user = _pick(t_cols, "Usuario_Creador", "Usuario")
    col_asig = _pick(t_cols, "Usuario_Asignado", "UsuarioAsignado", "Asignado_A", "Responsable")
    col_src = _pick(t_cols, "SourceType", "Source_Type")
    col_meta = _pick(t_cols, "Meta_JSON", "Datos_JSON", "Contexto_JSON")

    cols: List[str] = []
    vals: List[Any] = []
    if col_tipo:
        cols.append(col_tipo)
        vals.append("MANUAL")
    if col_titulo:
        cols.append(col_titulo)
        vals.append(titulo[:500])
    if col_desc:
        cols.append(col_desc)
        vals.append(desc[:4000])
    if col_estado:
        cols.append(col_estado)
        vals.append("Pendiente")
    if col_prog:
        cols.append(col_prog)
        vals.append(0)
    if col_user:
        cols.append(col_user)
        vals.append(payload.usuario[:120])
    if col_asig:
        cols.append(col_asig)
        vals.append(responsable)
    if col_src:
        cols.append(col_src)
        vals.append("Manual")
    if col_meta:
        cols.append(col_meta)
        vals.append(
            json.dumps(
                {
                    "_origen_crear_api": "MANUAL",
                    "categoria": "QA",
                    "id_reporte": id_reporte,
                    "usuarios_asignados": usuarios,
                },
                ensure_ascii=False,
            )
        )
    cur.execute(
        f"INSERT INTO Tbl_Gestor_Tareas ({', '.join(cols)}) OUTPUT INSERTED.{t_pk} VALUES ({', '.join(['?'] * len(vals))})",
        tuple(vals),
    )
    id_tarea = int(cur.fetchone()[0])

    if c_nombre:
        ccols = [c_fk, c_nombre]
        cvals: List[Any] = [id_tarea, f"Atender reporte QA #{id_reporte}"]
        if c_done:
            ccols.append(c_done)
            cvals.append(0)
        cur.execute(
            f"INSERT INTO Tbl_Gestor_Checklist ({', '.join(ccols)}) VALUES ({', '.join(['?'] * len(cvals))})",
            tuple(cvals),
        )

    if _tabla_existe(cur, "Tbl_Gestor_Tarea_Asignados"):
        for usr in usuarios:
            cur.execute(
                "INSERT INTO Tbl_Gestor_Tarea_Asignados (ID_Tarea, Usuario) VALUES (?, ?)",
                (id_tarea, usr),
            )
    return id_tarea


# === BUG TRACKER ===
@router.post("/api/reportes/nuevo")
def nuevo_reporte(payload: BugReportPayload):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        descripcion_full = payload.descripcion
        tags_norm: List[str] = []
        for raw_tag in (payload.hashtags or []):
            t = str(raw_tag or "").strip().lower().replace(" ", "_")
            if not t:
                continue
            if not t.startswith("#"):
                t = f"#{t}"
            if t not in tags_norm:
                tags_norm.append(t)
        if tags_norm:
            descripcion_full = f"{descripcion_full}\n\n[Tags]\n{' '.join(tags_norm)}"
        if payload.contexto_pantalla:
            descripcion_full = (
                f"{payload.descripcion}\n\n[Contexto pantalla]\n{payload.contexto_pantalla}"
            )
            if tags_norm:
                descripcion_full = f"{descripcion_full}\n\n[Tags]\n{' '.join(tags_norm)}"
        cursor.execute("""
            INSERT INTO Tbl_Reportes_Beta (Usuario, Fecha_Hora, Modulo, Descripcion, Gravedad, Estado, Captura_Base64)
            VALUES (?, GETDATE(), ?, ?, ?, 'Abierto', ?)
        """, (payload.usuario, payload.modulo, descripcion_full, payload.gravedad, payload.captura))
        id_reporte = None
        try:
            cursor.execute("SELECT CAST(SCOPE_IDENTITY() AS INT)")
            id_reporte = int(cursor.fetchone()[0])
        except Exception:
            id_reporte = None
        id_tarea = None
        if payload.crear_tarea_correccion and id_reporte is not None:
            id_tarea = _crear_tarea_correccion_ingenieria(cursor, id_reporte, payload)
        conn.commit()
        return {"status": "success", "id_reporte": id_reporte, "id_tarea": id_tarea}
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        conn.close()

@router.get("/api/reportes/exportar_gemini")
def exportar_reportes():
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT ID_Reporte, Modulo, Descripcion, Gravedad, Captura_Base64 FROM Tbl_Reportes_Beta WHERE Estado = 'Abierto'")
        rows = cursor.fetchall()
        reportes = []
        for r in rows:
            reportes.append({
                "id": r.ID_Reporte,
                "modulo": r.Modulo,
                "error": r.Descripcion,
                "severidad": r.Gravedad,
                "captura_base64": r.Captura_Base64
            })
        return reportes
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        conn.close()

@router.get("/api/reportes")
def listar_reportes():
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT ID_Reporte, Usuario, Fecha_Hora, Modulo, Descripcion, Gravedad, Estado, Captura_Base64 FROM Tbl_Reportes_Beta WHERE Estado = 'Abierto' ORDER BY Fecha_Hora DESC")
        rows = cursor.fetchall()
        reportes = []
        for r in rows:
            reportes.append({
                "id": r.ID_Reporte,
                "usuario": r.Usuario,
                "fecha": r.Fecha_Hora.strftime("%Y-%m-%d %H:%M:%S") if r.Fecha_Hora else None,
                "modulo": r.Modulo,
                "descripcion": r.Descripcion,
                "gravedad": r.Gravedad,
                "estado": r.Estado,
                "captura_base64": r.Captura_Base64
            })
        return reportes
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        conn.close()


@router.get("/api/reportes/historial")
def listar_reportes_historial():
    """Reportes ya gestionados: resueltos (Cerrado) o rechazados."""
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute(
            """
            SELECT ID_Reporte, Usuario, Fecha_Hora, Modulo, Descripcion, Gravedad, Estado, Captura_Base64
            FROM Tbl_Reportes_Beta
            WHERE Estado IN (N'Cerrado', N'Rechazado')
            ORDER BY Fecha_Hora DESC
            """
        )
        rows = cursor.fetchall()
        reportes = []
        for r in rows:
            reportes.append({
                "id": r.ID_Reporte,
                "usuario": r.Usuario,
                "fecha": r.Fecha_Hora.strftime("%Y-%m-%d %H:%M:%S") if r.Fecha_Hora else None,
                "modulo": r.Modulo,
                "descripcion": r.Descripcion,
                "gravedad": r.Gravedad,
                "estado": r.Estado,
                "captura_base64": r.Captura_Base64,
            })
        return reportes
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        conn.close()

@router.get("/api/reportes/exportar")
def exportar_reportes_excel():
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT ID_Reporte, Usuario, Fecha_Hora, Modulo, Descripcion, Gravedad, Estado FROM Tbl_Reportes_Beta ORDER BY Fecha_Hora DESC")
        rows = cursor.fetchall()
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Reportes Beta"
        
        headers = ["ID", "Usuario", "Fecha", "Módulo", "Gravedad", "Estado", "Descripción"]
        for col_idx, text in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=text)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
            cell.alignment = Alignment(horizontal="center")
            
        for row_idx, r in enumerate(rows, 2):
            fecha_str = r.Fecha_Hora.strftime("%Y-%m-%d %H:%M:%S") if r.Fecha_Hora else "Sin fecha"
            ws.cell(row=row_idx, column=1, value=int(r.ID_Reporte))
            ws.cell(row=row_idx, column=2, value=str(r.Usuario))
            ws.cell(row=row_idx, column=3, value=fecha_str)
            ws.cell(row=row_idx, column=4, value=str(r.Modulo))
            ws.cell(row=row_idx, column=5, value=str(r.Gravedad))
            ws.cell(row=row_idx, column=6, value=str(r.Estado))
            ws.cell(row=row_idx, column=7, value=str(r.Descripcion))

        # Auto-fit columns
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except: pass
            ws.column_dimensions[column].width = min((max_length + 2) * 1.1, 60)
            
        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)
        
        return StreamingResponse(
            stream,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": "attachment; filename=Reportes_Beta.xlsx"
            }
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        conn.close()

@router.put("/api/reportes/{id_reporte}/resolver")
def resolver_reporte(id_reporte: int):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute(
            "UPDATE Tbl_Reportes_Beta SET Estado = N'Cerrado' WHERE ID_Reporte = ?",
            (id_reporte,),
        )
        conn.commit()
        return {"status": "success"}
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        conn.close()


@router.put("/api/reportes/{id_reporte}/rechazar")
def rechazar_reporte(id_reporte: int):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute(
            "UPDATE Tbl_Reportes_Beta SET Estado = N'Rechazado' WHERE ID_Reporte = ?",
            (id_reporte,),
        )
        conn.commit()
        return {"status": "success"}
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        conn.close()


@router.delete("/api/reportes/limpiar_historial")
def limpiar_historial_reportes_beta(
    x_admin_master_password: Optional[str] = Header(None, alias="X-Admin-Master-Password"),
):
    """Elimina filas con Estado Cerrado o Rechazado (reportes ya gestionados)."""
    assert_admin_master_password_matches(x_admin_master_password)
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute(
            """
            DELETE FROM Tbl_Reportes_Beta
            WHERE Estado = N'Cerrado' OR Estado = N'Rechazado'
            """
        )
        n = int(cursor.rowcount or 0)
        registrar_log_global(
            cursor,
            "QA_REPORTES",
            "LIMPIAR_HISTORIAL",
            "",
            f"borradas={n}",
            "admin_master",
        )
        conn.commit()
        return {"status": "ok", "borradas": n}
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        conn.close()
